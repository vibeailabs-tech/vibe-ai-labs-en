import sys; sys.path.insert(0,"/tmp/brandbuild")
from cfg import *
import tpl, gfx, logo, render
from tpl import T, grad_def, wrap_svg, C, HF, BF, icon_g, pill
W,H=1920,1080
def slide_head(dark,n=None,uidp="p"):
    fg="#FFFFFF" if dark else C["midnight"]
    lg=logo.horizontal_parts("reversed" if dark else "color","h"+tpl.uid(),False)
    ld,lb,lw,lh=logo.place(lg,110,64,height=46)
    foot=T(110,H-52,"Vibe AI Labs · "+BRAND_LINE,20,500,"#D9E2F2" if dark else C["slate"],BF)
    if n: foot+=T(W-110,H-52,"%02d"%n,20,700,"#D9E2F2" if dark else C["slate"],BF,"end")
    return ld,lb+foot
def S(name,dark,build,n=None,bgk="ribbon",bgp=None,head=True):
    if bgk=="ribbon": d,b=gfx.ribbon(W,H,dark,uid="s"+tpl.uid(),**(bgp or dict(cx=0.86,cy=0.62,wr=0.95)))
    elif bgk=="waves": d,b=gfx.waves(W,H,dark,n=40,amp=0.12,y0=0.82,uid="s"+tpl.uid())
    else: d,b=gfx.glows(W,H,dark,"s"+tpl.uid()).replace("<defs>","<defs>",1),""; d,b="",gfx.glows(W,H,dark,"s"+tpl.uid())
    defs=d; body=b
    if head:
        hd,hb=slide_head(dark,n); defs+=hd; body+=hb
    dd,bb=build(dark); defs+=dd; body+=bb
    svg=wrap_svg(W,H,defs,body)
    w("06_Templates/Presentation/SVG/%s.svg"%name,svg); jobs.append((svg,os.path.join(OUT,"06_Templates/Presentation/PNG/%s.png"%name),W,H,False))
jobs=[]
fgc=lambda dk:"#FFFFFF" if dk else C["midnight"]
subc=lambda dk:"#D9E2F2" if dk else C["slate"]
def gradfill(dk,x,w_=900):
    gid=tpl.uid("g"); return gid, grad_def(gid,x,0,x+w_,0,"#5B8CFF" if dk else "#2563FF","#14D8C4")

def cover(dk):
    gid,gd=gradfill(dk,110,1100)
    b=T(110,430,"BÀI THUYẾT TRÌNH",24,700,C["aqua"] if dk else C["blue"],HF,ls=6)
    b+=T(110,560,"Tên bài thuyết trình",110,800,fgc(dk),HF,ls=-2)+T(110,690,"hai dòng tối đa.",110,800,"url(#%s)"%gid,HF,ls=-2)
    b+=T(110,780,"Tiêu đề phụ — Marketing × AI × Creative",34,500,subc(dk),BF)+T(110,960,"[Người trình bày] · [Tháng / Năm]",24,500,subc(dk),BF)
    return gd,b
def agenda(dk):
    b=T(110,250,"NỘI DUNG",22,700,C["blue"],HF,ls=6)+T(110,350,"Chương trình",72,800,C["midnight"],HF,ls=-1)
    items=["Bối cảnh & bài toán","Chiến lược & giải pháp","Lộ trình triển khai","Đo lường & bước tiếp theo"]
    for i,t in enumerate(items):
        y=470+i*128
        b+='<rect x="110" y="%d" width="1180" height="100" rx="24" fill="#FFFFFF" stroke="#D9E2F2"/>'%y+T(160,y+66,"%02d"%(i+1),40,800,C["blue"],HF)+T(260,y+64,t,36,700,C["midnight"],HF)
    return "",b
def section(dk):
    gid,gd=gradfill(dk,110,700)
    b=T(110,560,"01",300,800,"url(#%s)"%gid,HF,ls=-8)+T(110,720,"Tên phần / chương",84,800,fgc(dk),HF,ls=-1)+T(110,790,"Một câu mô tả ngắn cho phần này.",32,500,subc(dk),BF)
    return gd,b
def pillars(dk):
    b=T(110,250,"HỆ THỐNG TĂNG TRƯỞNG",22,700,C["blue"],HF,ls=6)+T(110,340,"Ba năng lực. Một hệ thống.",68,800,C["midnight"],HF,ls=-1)
    cols=[("growth","Marketing số","Chiến lược · Thương hiệu · Nội dung · Mạng xã hội · Website · Performance · SEO · PR · Dữ liệu · Thương mại điện tử",C["blue"]),
          ("ai","AI & Tự động hóa","Tư vấn AI · AI Agent · Chatbot & trợ lý · AI tạo sinh · Tích hợp & tự động hóa quy trình",C["green"]),
          ("design","Thiết kế & Sáng tạo","Nhận diện · Đồ họa · Minh họa · Nhiếp ảnh & video · Sản xuất bằng AI tạo sinh",C["orange"])]
    cw=540; gap=40
    for i,(ic,t,s_,col) in enumerate(cols):
        x=110+i*(cw+gap)
        b+='<rect x="%d" y="440" width="%d" height="520" rx="36" fill="#FFFFFF" stroke="#D9E2F2"/>'%(x,cw)
        b+='<rect x="%d" y="480" width="112" height="112" rx="30" fill="%s" fill-opacity=".14"/>'%(x+40,col)+icon_g(ic,x+40+24,480+24,64,col,1.75)
        b+=T(x+40,670,t,42,800,C["midnight"],HF)+'<rect x="%d" y="694" width="64" height="5" rx="2.5" fill="%s"/>'%(x+40,col)
        words=s_.split(" · "); lines=[]; cur=""
        for wd in words:
            if len(cur)+len(wd)>30: lines.append(cur.rstrip(" ·")); cur=wd+" · "
            else: cur+=wd+" · "
        lines.append(cur.rstrip(" ·"))
        for j,l in enumerate(lines[:5]): b+=T(x+40,750+j*40,l,25,500,C["slate"],BF)
    return "",b
def process(dk):
    b=T(110,250,"PHƯƠNG PHÁP",22,700,C["aqua"],HF,ls=6)+T(110,350,"Audit. Build. Automate. Grow.",72,800,"#FFFFFF",HF,ls=-1)
    steps=[("audit","Audit","Đánh giá hiện trạng: thị trường, khách hàng, dữ liệu, quy trình."),("design","Build","Xây nền tảng: chiến lược, thương hiệu, nội dung, đo lường."),
           ("automation","Automate","Kết nối hệ thống, AI Agent, chatbot và quy trình tự động."),("growth","Grow","Thử nghiệm, đo lường và tối ưu theo dữ liệu thực tế.")]
    cw=410; gap=26
    for i,(ic,t,s_) in enumerate(steps):
        x=110+i*(cw+gap)
        b+='<rect x="%d" y="470" width="%d" height="440" rx="32" fill="#FFFFFF" fill-opacity=".07" stroke="#FFFFFF" stroke-opacity=".22"/>'%(x,cw)
        b+=T(x+36,540,"%02d"%(i+1),30,800,C["aqua"],HF)+icon_g(ic,x+cw-36-56,500,56,"#FFFFFF",1.75)+T(x+36,650,t,50,800,"#FFFFFF",HF)
        ws=s_.split(" "); lines=[]; cur=""
        for wd in ws:
            if len(cur)+len(wd)>26: lines.append(cur.strip()); cur=wd+" "
            else: cur+=wd+" "
        lines.append(cur.strip())
        for j,l in enumerate(lines): b+=T(x+36,720+j*38,l,25,500,"#D9E2F2",BF)
    return "",b
def content(dk):
    b=T(110,250,"CHỦ ĐỀ",22,700,C["blue"],HF,ls=6)+T(110,350,"Tiêu đề nội dung chính",64,800,C["midnight"],HF,ls=-1)
    pts=["Điểm chính thứ nhất — một câu ngắn, có động từ cụ thể.","Điểm chính thứ hai — dẫn bằng bằng chứng hoặc số liệu.","Điểm chính thứ ba — kết thúc bằng hành động tiếp theo."]
    for i,t in enumerate(pts):
        y=470+i*120
        b+='<circle cx="130" cy="%d" r="14" fill="#2563FF"/>'%(y-10)+T(170,y,t,34,500,C["ink"],BF)
    b+='<rect x="1100" y="260" width="710" height="640" rx="36" fill="#E9F0FF" stroke="#D9E2F2"/>'+T(1455,590,"Khu vực hình ảnh / biểu đồ",30,600,C["slate"],BF,"middle")
    return "",b
def stats(dk):
    gid,gd=gradfill(dk,110,1700)
    b=T(110,250,"KẾT QUẢ",22,700,C["aqua"],HF,ls=6)+T(110,350,"Con số nói lên điều gì",68,800,"#FFFFFF",HF,ls=-1)
    for i,(n,l) in enumerate((("+00%","Chỉ số thứ nhất"),("00×","Chỉ số thứ hai"),("−00%","Chỉ số thứ ba"))):
        x=110+i*590
        b+=T(x,640,n,150,800,"url(#%s)"%gid,HF,ls=-4)+T(x,720,l,34,700,"#FFFFFF",HF)+T(x,766,"Nguồn · thời gian đo (chỉ dùng số liệu xác minh)",22,500,"#D9E2F2",BF)
    return gd,b
def quote(dk):
    gid,gd=gradfill(dk,110,400)
    b=T(110,420,"“",340,800,"url(#%s)"%gid,HF)+T(220,500,"Nội dung nhận xét của khách hàng đặt tại đây —",56,700,"#FFFFFF",HF)+T(220,580,"ngắn gọn, cụ thể và có kết quả.",56,700,"#FFFFFF",HF)
    b+=T(220,720,"Tên khách hàng",34,800,"#FFFFFF",HF)+T(220,766,"Chức vụ · Doanh nghiệp",26,500,"#D9E2F2",BF)
    return gd,b
def closing(dk):
    gid,gd=gradfill(dk,110,1100)
    b=T(110,420,"Cùng xây dựng",110,800,"#FFFFFF",HF,ls=-2)+T(110,540,"điều tiếp theo.",110,800,"url(#%s)"%gid,HF,ls=-2)
    b+=T(110,660,"Trao đổi với Vibe AI Labs để xác định bài toán và xây dựng giải pháp phù hợp.",32,500,"#D9E2F2",BF)
    yy=760
    for t in (EMAIL,WEBSITE_SHORT,ADDRESS1+", "+ADDRESS2):
        b+=T(110,yy,t,30,600,"#FFFFFF",BF); yy+=50
    return gd,b
S("01-cover_dark",True,cover,bgp=dict(cx=0.86,cy=0.64,wr=1.0))
S("02-agenda_light",False,agenda,2,bgk="plain")
S("03-section_dark",True,section,3,bgp=dict(cx=0.9,cy=0.7,wr=0.8))
S("04-pillars_light",False,pillars,4,bgk="plain")
S("05-process_dark",True,process,5,bgk="waves")
S("06-content_light",False,content,6,bgk="plain")
S("07-results_dark",True,stats,7,bgk="waves")
S("08-quote_dark",True,quote,8,bgp=dict(cx=0.92,cy=0.72,wr=0.8))
S("09-closing_dark",True,closing,9,bgp=dict(cx=0.84,cy=0.62,wr=1.0))
res=render.batch(jobs,workers=6); print("png ok",sum(res),len(res))

# ---------------------------------------------------------------- PPTX (editable)
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
def rgb(h): h=h.lstrip("#"); return RGBColor(int(h[:2],16),int(h[2:4],16),int(h[4:],16))
prs=Presentation(); prs.slide_width=Emu(12192000); prs.slide_height=Emu(6858000)
PX=12192000/1920.0
def E(v): return Emu(int(v*PX))
def tb(sl,x,y,wd,ht,text,size,color,font=HF,bold=True,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP):
    t=sl.shapes.add_textbox(E(x),E(y),E(wd),E(ht)); tf=t.text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    lines=text if isinstance(text,list) else [text]
    for i,l in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.alignment=align
        r=p.add_run(); r.text=l; r.font.size=Pt(size*0.75); r.font.bold=bold; r.font.name=font; r.font.color.rgb=rgb(color)
    return t
def box(sl,x,y,wd,ht,fill,line=None,radius=0.1,alpha=None):
    sh=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,E(x),E(y),E(wd),E(ht)); sh.adjustments[0]=radius
    sh.fill.solid(); sh.fill.fore_color.rgb=rgb(fill)
    if line: sh.line.color.rgb=rgb(line); sh.line.width=Pt(0.75)
    else: sh.line.fill.background()
    sh.shadow.inherit=False; return sh
base=OUT+"/04_Graphic_Elements/PNG/"
logo_png=lambda dk: OUT+"/01_Logos/PNG/vibe-ai-labs_horizontal_%s_1000px.png"%("reversed" if dk else "color")
def new(bg_file,dk,n=None):
    sl=prs.slides.add_slide(prs.slide_layouts[6])
    if bg_file: sl.shapes.add_picture(bg_file,0,0,prs.slide_width,prs.slide_height)
    else:
        sl.background.fill.solid(); sl.background.fill.fore_color.rgb=rgb("#F8FAFC")
    sl.shapes.add_picture(logo_png(dk),E(110),E(64),height=E(46))
    tb(sl,110,H-70,1200,30,"Vibe AI Labs · "+BRAND_LINE,20,"#D9E2F2" if dk else C["slate"],BF,False)
    if n: tb(sl,W-310,H-70,200,30,"%02d"%n,20,"#D9E2F2" if dk else C["slate"],BF,True,PP_ALIGN.RIGHT)
    return sl
D="dark"; Lt="light"
rb=lambda tone: base+"ribbon-background_%s_16x9.png"%tone
wv=lambda tone: base+"wave-lines_%s_16x9.png"%tone
s=new(rb(D),True); tb(s,110,400,900,30,"BÀI THUYẾT TRÌNH",24,C["aqua"]); tb(s,110,440,1300,300,["Tên bài thuyết trình","hai dòng tối đa."],104,"#FFFFFF"); tb(s,110,760,1200,50,"Tiêu đề phụ — Marketing × AI × Creative",34,"#D9E2F2",BF,False); tb(s,110,940,1200,40,"[Người trình bày] · [Tháng / Năm]",24,"#D9E2F2",BF,False)
s=new(None,False,2); tb(s,110,230,600,30,"NỘI DUNG",22,C["blue"]); tb(s,110,290,1200,100,"Chương trình",72,C["midnight"])
for i,t in enumerate(["Bối cảnh & bài toán","Chiến lược & giải pháp","Lộ trình triển khai","Đo lường & bước tiếp theo"]):
    y=470+i*128; box(s,110,y,1180,100,"#FFFFFF",C["line"],0.25); tb(s,160,y,100,100,"%02d"%(i+1),40,C["blue"],HF,True,PP_ALIGN.LEFT,MSO_ANCHOR.MIDDLE); tb(s,260,y,900,100,t,36,C["midnight"],HF,True,PP_ALIGN.LEFT,MSO_ANCHOR.MIDDLE)
s=new(rb(D),True,3); tb(s,110,300,900,360,"01",300,C["aqua"]); tb(s,110,640,1400,120,"Tên phần / chương",84,"#FFFFFF"); tb(s,110,770,1400,50,"Một câu mô tả ngắn cho phần này.",32,"#D9E2F2",BF,False)
s=new(None,False,4); tb(s,110,230,800,30,"HỆ THỐNG TĂNG TRƯỞNG",22,C["blue"]); tb(s,110,280,1500,100,"Ba năng lực. Một hệ thống.",68,C["midnight"])
for i,(t,sub,col) in enumerate((("Marketing số","Chiến lược · Thương hiệu · Nội dung · Mạng xã hội · Website · Performance · SEO · PR · Dữ liệu · Thương mại điện tử",C["blue"]),("AI & Tự động hóa","Tư vấn AI · AI Agent · Chatbot & trợ lý · AI tạo sinh · Tích hợp & tự động hóa",C["green"]),("Thiết kế & Sáng tạo","Nhận diện · Đồ họa · Minh họa · Nhiếp ảnh & video · Sản xuất bằng AI tạo sinh",C["orange"]))):
    x=110+i*580; box(s,x,440,540,520,"#FFFFFF",C["line"],0.07); b2=box(s,x+40,480,112,112,col,None,0.27); b2.fill.fore_color.rgb=rgb(col)
    tb(s,x+40,640,460,60,t,42,C["midnight"]); tb(s,x+40,730,460,200,sub,25,C["slate"],BF,False)
s=new(wv(D),True,5); tb(s,110,230,800,30,"PHƯƠNG PHÁP",22,C["aqua"]); tb(s,110,280,1700,100,"Audit. Build. Automate. Grow.",72,"#FFFFFF")
for i,(t,sub) in enumerate((("Audit","Đánh giá hiện trạng: thị trường, khách hàng, dữ liệu, quy trình."),("Build","Xây nền tảng: chiến lược, thương hiệu, nội dung, đo lường."),("Automate","Kết nối hệ thống, AI Agent, chatbot và quy trình tự động."),("Grow","Thử nghiệm, đo lường và tối ưu theo dữ liệu thực tế."))):
    x=110+i*436; sh=box(s,x,470,410,440,"#1B2347","#3A4572",0.08); tb(s,x+36,510,100,40,"%02d"%(i+1),30,C["aqua"]); tb(s,x+36,610,340,70,t,50,"#FFFFFF"); tb(s,x+36,700,340,180,sub,25,"#D9E2F2",BF,False)
s=new(None,False,6); tb(s,110,230,600,30,"CHỦ ĐỀ",22,C["blue"]); tb(s,110,280,1500,90,"Tiêu đề nội dung chính",64,C["midnight"])
tb(s,110,450,900,400,["•  Điểm chính thứ nhất — một câu ngắn, có động từ cụ thể.","•  Điểm chính thứ hai — dẫn bằng bằng chứng hoặc số liệu.","•  Điểm chính thứ ba — kết thúc bằng hành động tiếp theo."],34,C["ink"],BF,False)
box(s,1100,260,710,640,"#E9F0FF",C["line"],0.06); tb(s,1100,560,710,60,"Khu vực hình ảnh / biểu đồ",30,C["slate"],BF,False,PP_ALIGN.CENTER)
s=new(wv(D),True,7); tb(s,110,230,600,30,"KẾT QUẢ",22,C["aqua"]); tb(s,110,280,1500,100,"Con số nói lên điều gì",68,"#FFFFFF")
for i,(n,l) in enumerate((("+00%","Chỉ số thứ nhất"),("00×","Chỉ số thứ hai"),("−00%","Chỉ số thứ ba"))):
    x=110+i*590; tb(s,x,480,560,200,n,160,C["aqua"]); tb(s,x,690,560,50,l,34,"#FFFFFF"); tb(s,x,745,560,40,"Nguồn · thời gian đo (chỉ dùng số liệu xác minh)",22,"#D9E2F2",BF,False)
s=new(rb(D),True,8); tb(s,110,260,300,300,"“",300,C["aqua"]); tb(s,220,440,1300,200,["Nội dung nhận xét của khách hàng đặt tại đây —","ngắn gọn, cụ thể và có kết quả."],56,"#FFFFFF"); tb(s,220,690,900,50,"Tên khách hàng",34,"#FFFFFF"); tb(s,220,740,900,40,"Chức vụ · Doanh nghiệp",26,"#D9E2F2",BF,False)
s=new(rb(D),True,9); tb(s,110,300,1500,300,["Cùng xây dựng","điều tiếp theo."],104,"#FFFFFF"); tb(s,110,630,1400,50,"Trao đổi với Vibe AI Labs để xác định bài toán và xây dựng giải pháp phù hợp.",32,"#D9E2F2",BF,False); tb(s,110,730,1400,160,[EMAIL,WEBSITE_SHORT,ADDRESS1+", "+ADDRESS2],30,"#FFFFFF",BF,True)
prs.save(os.path.join(OUT,"06_Templates/Presentation/Vibe_AI_Labs_Presentation_Template.pptx")); print("pptx saved")
