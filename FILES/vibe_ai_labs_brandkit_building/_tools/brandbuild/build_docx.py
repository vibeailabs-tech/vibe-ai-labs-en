import sys; sys.path.insert(0,"/tmp/brandbuild")
from cfg import *
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
def rgb(h): h=h.lstrip("#"); return RGBColor(int(h[:2],16),int(h[2:4],16),int(h[4:],16))
d=Document()
sec=d.sections[0]; sec.page_width=Cm(21); sec.page_height=Cm(29.7); sec.left_margin=sec.right_margin=Cm(2.2); sec.top_margin=Cm(3.2); sec.bottom_margin=Cm(2.6)
def setfont(style,name,size,color,bold=False):
    style.font.name=name; style.font.size=Pt(size); style.font.bold=bold; style.font.color.rgb=rgb(color)
    rpr=style.element.get_or_add_rPr(); rf=rpr.find(qn("w:rFonts"))
    if rf is None: rf=OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii","w:hAnsi","w:cs","w:eastAsia"): rf.set(qn(a),name)
st=d.styles
setfont(st["Normal"],"Inter",10.5,C["ink"]); st["Normal"].paragraph_format.space_after=Pt(6); st["Normal"].paragraph_format.line_spacing=1.35
setfont(st["Title"],"Plus Jakarta Sans",30,C["midnight"],True); st["Title"].paragraph_format.space_after=Pt(4)
for n,(sz,col) in {"Heading 1":(20,C["midnight"]),"Heading 2":(15,C["blue"]),"Heading 3":(12,C["midnight"])}.items():
    setfont(st[n],"Plus Jakarta Sans",sz,col,True); st[n].paragraph_format.space_before=Pt(14 if n=="Heading 1" else 10); st[n].paragraph_format.space_after=Pt(4)
setfont(st["List Bullet"],"Inter",10.5,C["ink"])
# header with logo
hdr=sec.header; hp=hdr.paragraphs[0]; hp.add_run().add_picture(OUT+"/01_Logos/PNG/vibe-ai-labs_horizontal_color_1000px.png",width=Cm(4.6))
pPr=hp._p.get_or_add_pPr(); b=OxmlElement("w:pBdr"); bt=OxmlElement("w:bottom"); bt.set(qn("w:val"),"single"); bt.set(qn("w:sz"),"6"); bt.set(qn("w:space"),"6"); bt.set(qn("w:color"),"D9E2F2"); b.append(bt); pPr.append(b)
# footer
fp=sec.footer.paragraphs[0]; fp.text=""
r=fp.add_run("Vibe AI Labs · %s · %s   |   "%(EMAIL,WEBSITE_SHORT)); r.font.size=Pt(8); r.font.color.rgb=rgb(C["slate"]); r.font.name="Inter"
def field(par,code):
    run=par.add_run(); run.font.size=Pt(8); run.font.color.rgb=rgb(C["slate"])
    a=OxmlElement("w:fldChar"); a.set(qn("w:fldCharType"),"begin"); i=OxmlElement("w:instrText"); i.set(qn("xml:space"),"preserve"); i.text=code; e=OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"),"end")
    run._r.append(a); run._r.append(i); run._r.append(e)
field(fp,"PAGE")
# content
p=d.add_paragraph("PROPOSAL · BÁO CÁO · TÀI LIỆU"); p.runs[0].font.size=Pt(9); p.runs[0].font.bold=True; p.runs[0].font.color.rgb=rgb(C["blue"]); p.runs[0].font.name="Plus Jakarta Sans"
d.add_paragraph("Tiêu đề tài liệu",style="Title")
q=d.add_paragraph("Phụ đề ngắn — Marketing × AI × Creative"); q.runs[0].font.size=Pt(13); q.runs[0].font.color.rgb=rgb(C["slate"])
d.add_paragraph("Đoạn mở đầu: nói rõ bài toán hoặc lợi ích trong câu đầu tiên. Dùng câu ngắn, động từ cụ thể; tránh phóng đại. Thân bài dùng Inter 10,5 pt, giãn dòng 1,35.")
d.add_heading("1. Tiêu đề cấp 1",level=1)
d.add_paragraph("Nội dung thân bài. Số liệu chỉ dùng khi có thể xác minh, kèm nguồn và kỳ đo.")
d.add_heading("1.1 Tiêu đề cấp 2",level=2)
for t in ("Gạch đầu dòng ngắn, bắt đầu bằng động từ.","Mỗi ý một dòng; tối đa 5–7 ý.","Kết thúc bằng hành động tiếp theo."): d.add_paragraph(t,style="List Bullet")
d.add_heading("1.2 Bảng mẫu",level=2)
t=d.add_table(rows=4,cols=3); t.style="Light Grid Accent 1"
hdrs=("Hạng mục","Mô tả","Thời gian"); 
for i,h in enumerate(hdrs): t.rows[0].cells[i].text=h
rows=[("Audit","Đánh giá hiện trạng, dữ liệu, quy trình","Tuần 1–2"),("Build","Xây nền tảng, nội dung, đo lường","Tuần 3–6"),("Automate / Grow","Tự động hóa, thử nghiệm, tối ưu","Tuần 7–12")]
for r_,row in enumerate(rows,1):
    for c_,v in enumerate(row): t.rows[r_].cells[c_].text=v
for row in t.rows:
    for c in row.cells:
        for pp in c.paragraphs:
            for rr in pp.runs: rr.font.size=Pt(9.5); rr.font.name="Inter"
d.add_paragraph()
cp=d.add_paragraph("Ghi chú: sử dụng phong cách tiêu đề (Heading 1–3) thay vì chỉnh tay để mục lục tự động hoạt động. Font Plus Jakarta Sans và Inter nằm trong 03_Typography/fonts."); cp.runs[0].font.size=Pt(9); cp.runs[0].font.color.rgb=rgb(C["slate"]); cp.runs[0].font.italic=True
os.makedirs(OUT+"/06_Templates/Documents",exist_ok=True)
d.save(OUT+"/06_Templates/Documents/Vibe_AI_Labs_Document_Template.docx"); print("docx ok")
